"""Report builder tests (plan sections 7, 8 and 13).

Builds the three reports from a freshly generated seed-42 database (shared `built`
fixture in conftest.py) into a temporary folder, then checks the HTML, numbers,
Excel file and charts. Also checks the planted Apparel decline is visible.
"""

import html
import io
import re
from datetime import date
from pathlib import Path

import matplotlib.image as mpimg
import openpyxl
import pandas as pd
import pytest
import yaml

from src.reports import builder
from src.reports.builder import Report

AS_OF = {
    "daily_sales_snapshot": date(2026, 9, 29),
    "weekly_sales_review": date(2026, 9, 27),
    "monthly_business_review": date(2026, 8, 31),
}
EXCEL_SHEETS = [
    "KPIs",
    "Categories",
    "Channels",
    "Regions",
    "Customer types",
    "Top products",
    "Top customers",
]
PNG_SIGNATURE = b"\x89PNG\r\n\x1a\n"
EXPECTED_CHART_WIDTH_PX = 1200  # 6 inches at 200 dpi
NUMBER = re.compile(r"\d[\d,]*(?:\.\d+)?")
TAG = re.compile(r"<[^>]+>")


@pytest.fixture(scope="module")
def reports(built: dict, tmp_path_factory: pytest.TempPathFactory) -> dict[str, Report]:
    """All three reports built once from the generated database."""
    out = tmp_path_factory.mktemp("reports")
    return {
        name: builder.build_report(name, as_of, db_path=built["db_path"], output_root=out)
        for name, as_of in AS_OF.items()
    }


def visible_text(page: str) -> str:
    """The text a reader sees: tags removed, entities decoded."""
    return html.unescape(TAG.sub(" ", page))


def leaves(value: object) -> list[object]:
    """Every scalar inside a nested dict/list."""
    if isinstance(value, dict):
        return [leaf for v in value.values() for leaf in leaves(v)]
    if isinstance(value, list):
        return [leaf for v in value for leaf in leaves(v)]
    return [value]


def allowed_numbers(numbers: dict) -> set[str]:
    """Every numeric token the HTML may show, derived from Report.numbers only.

    Numbers are allowed in any of the builder's formats; text values (labels, dates,
    section titles) contribute the digits they contain.
    """
    allowed: set[str] = set()
    for leaf in leaves(numbers):
        if isinstance(leaf, bool) or leaf is None:
            continue
        if isinstance(leaf, int | float):
            shown = [
                builder.fmt_money(leaf),
                builder.fmt_count(leaf),
                builder.fmt_pct(leaf),
                builder.fmt_change(leaf, "pct"),
                builder.fmt_change(leaf, "pp"),
                builder.fmt_days(leaf),
            ]
        else:
            shown = [str(leaf)]
        for text in shown:
            allowed.update(NUMBER.findall(text))
    # Fixed wording outside the numbers: the Week 4 summary placeholder.
    allowed.update(NUMBER.findall(builder.SUMMARY_PLACEHOLDER))
    return allowed


@pytest.mark.parametrize("name", list(AS_OF))
def test_no_template_tags_left(reports: dict[str, Report], name: str) -> None:
    page = reports[name].html
    for tag in ("{{", "}}", "{%", "%}", "{#", "#}"):
        assert tag not in page


@pytest.mark.parametrize("name", list(AS_OF))
def test_every_number_in_html_comes_from_numbers(reports: dict[str, Report], name: str) -> None:
    report = reports[name]
    shown = set(NUMBER.findall(visible_text(report.html)))
    unexplained = shown - allowed_numbers(report.numbers)
    assert not unexplained, f"numbers in HTML not in Report.numbers: {sorted(unexplained)}"


@pytest.mark.parametrize("name", list(AS_OF))
def test_headline_net_revenue_is_shown(reports: dict[str, Report], name: str) -> None:
    report = reports[name]
    kpi_section = next(s for s in report.numbers["sections"].values() if s["type"] == "kpis")
    net = next(r for r in kpi_section["rows"] if r["kpi"] == "net_revenue")
    assert net["current"] > 0
    assert builder.fmt_money(net["current"]) in visible_text(report.html)


@pytest.mark.parametrize("name", list(AS_OF))
def test_placeholders_present(reports: dict[str, Report], name: str) -> None:
    text = visible_text(reports[name].html)
    assert builder.SUMMARY_PLACEHOLDER in text
    assert "Added in Week 5" not in text  # the Week 2 placeholders are filled


def test_demand_sections_by_report(reports: dict[str, Report]) -> None:
    sections = {name: set(r.numbers["sections"]) for name, r in reports.items()}
    assert "demand_alerts" not in sections["daily_sales_snapshot"]
    assert "forecast_accuracy" not in sections["daily_sales_snapshot"]
    assert "demand_alerts" in sections["weekly_sales_review"]
    assert {"demand_alerts", "forecast_accuracy"} <= sections["monthly_business_review"]


def test_demand_lists_are_cut_at_10_rows_with_and_n_more(reports: dict[str, Report]) -> None:
    report = reports["weekly_sales_review"]
    overdue = next(
        g for g in report.numbers["sections"]["demand_alerts"]["groups"] if g["key"] == "overdue"
    )
    assert overdue["count"] > 10 and len(overdue["rows"]) == overdue["count"]  # all stored
    text = visible_text(report.html)
    assert overdue["more"] == overdue["count"] - 10
    assert f"and {builder.fmt_count(overdue['more'])} more" in text
    sheet = openpyxl.load_workbook(report.excel_path)["Demand alerts"]
    lists = [row[0] for row in sheet.iter_rows(min_row=2, values_only=True)]
    assert lists.count(overdue["title"]) == overdue["count"]


def test_daily_has_no_charts_or_excel(reports: dict[str, Report]) -> None:
    report = reports["daily_sales_snapshot"]
    assert report.images == {}
    assert report.excel_path is None


@pytest.mark.parametrize("name", ["weekly_sales_review", "monthly_business_review"])
def test_excel_sheets_and_formats(reports: dict[str, Report], name: str) -> None:
    path = reports[name].excel_path
    assert path is not None and path.exists()
    workbook = openpyxl.load_workbook(path)
    demand_sheets = {
        "weekly_sales_review": ["Demand alerts"],
        "monthly_business_review": ["Demand alerts", "Forecast accuracy"],
    }
    assert workbook.sheetnames == EXCEL_SHEETS + demand_sheets[name]
    for sheet in workbook.worksheets:
        assert sheet.freeze_panes == "A2"
    categories = workbook["Categories"]
    headers = [c.value for c in categories[1]]
    assert categories.cell(2, headers.index("Net revenue") + 1).number_format == "#,##0.00"
    assert categories.cell(2, headers.index("Share") + 1).number_format == "0.0%"


@pytest.mark.parametrize(
    ("name", "charts"),
    [
        ("weekly_sales_review", {"revenue_trend_13w", "revenue_by_category"}),
        ("monthly_business_review", {"revenue_trend_12m", "margin_by_category"}),
    ],
)
def test_charts_are_valid_pngs(reports: dict[str, Report], name: str, charts: set[str]) -> None:
    report = reports[name]
    assert set(report.images) == charts
    for cid, png in report.images.items():
        assert png.startswith(PNG_SIGNATURE)
        assert mpimg.imread(io.BytesIO(png), format="png").shape[1] == EXPECTED_CHART_WIDTH_PX
        assert f'src="cid:{cid}"' in report.html


@pytest.mark.parametrize("name", list(AS_OF))
def test_dry_run_files_open_in_a_browser(reports: dict[str, Report], name: str) -> None:
    report = reports[name]
    saved = (report.output_dir / builder.HTML_FILENAME).read_text(encoding="utf-8")
    assert "cid:" not in saved
    for cid in report.images:
        assert (report.output_dir / f"{cid}.png").exists()


def test_planted_apparel_decline_visible_in_weekly(reports: dict[str, Report]) -> None:
    report = reports["weekly_sales_review"]
    rows = report.numbers["sections"]["category_breakdown"]["rows"]
    apparel = next(r for r in rows if r["value"] == "Apparel")
    assert apparel["net_revenue_change"] < 0
    assert builder.fmt_change(apparel["net_revenue_change"], "pct").startswith("-")
    assert builder.fmt_change(apparel["net_revenue_change"], "pct") in report.html


def kpi_change(report: Report, section: str, kpi: str, comparison: str) -> float:
    """The stored change of one KPI against one comparison."""
    rows = report.numbers["sections"][section]["rows"]
    return next(r for r in rows if r["kpi"] == kpi)["comparisons"][comparison]["change"]


def test_percent_change_is_rounded_once(reports: dict[str, Report]) -> None:
    # Regression (Week 4 wrap-up): August 2026 net revenue vs August 2025 is
    # +8.6495%. Rounding it to 2 decimals first (8.65) and then to 1 showed "+8.7%".
    report = reports["monthly_business_review"]
    change = kpi_change(report, "kpis_month", "net_revenue", "same_month_last_year")
    assert change == pytest.approx(8.6495, abs=0.0001)  # stored unrounded
    assert builder.fmt_change(change, "pct") == "+8.6%"
    row = html.unescape(report.html)
    assert "+8.6%" in row and "+8.7%" not in row


def test_apparel_change_matches_hand_check(reports: dict[str, Report]) -> None:
    # Raw -25.346%: rounded once it is -25.3% (rounded twice it showed -25.4%).
    report = reports["weekly_sales_review"]
    rows = report.numbers["sections"]["category_breakdown"]["rows"]
    apparel = next(r for r in rows if r["value"] == "Apparel")
    assert builder.fmt_change(apparel["net_revenue_change"], "pct") == "-25.3%"


def test_rates_are_stored_unrounded(reports: dict[str, Report]) -> None:
    rows = reports["weekly_sales_review"].numbers["sections"]["category_breakdown"]["rows"]
    shares = [r["share"] for r in rows]
    assert any(round(share, 4) != share for share in shares)


@pytest.mark.parametrize(
    ("name", "as_of"),
    [
        ("weekly_sales_review", date(2026, 9, 26)),  # a Saturday
        ("monthly_business_review", date(2026, 9, 29)),  # not a month end
        ("weekly_sales_review", date(2026, 10, 4)),  # after the last order
        ("no_such_report", date(2026, 9, 27)),
    ],
)
def test_bad_requests_are_rejected(built: dict, tmp_path: Path, name: str, as_of: date) -> None:
    with pytest.raises(ValueError):
        builder.build_report(name, as_of, db_path=built["db_path"], output_root=tmp_path)


def test_unknown_query_name_in_config_is_rejected(tmp_path: Path) -> None:
    config = yaml.safe_load(builder.CONFIG_PATH.read_text(encoding="utf-8"))
    config["weekly_sales_review"]["queries"].append("top_widgets")
    path = tmp_path / "reports.yaml"
    path.write_text(yaml.safe_dump(config), encoding="utf-8")
    with pytest.raises(ValueError, match="top_widgets"):
        builder.load_config(path)


def test_money_is_rounded_half_up_from_the_exact_sum(reports: dict[str, Report]) -> None:
    # Regression (Week 5): August 2026 net revenue is exactly 462,431.555 in SQL.
    # Python's round() on the float gave 462,431.55; SQL ROUND gives 462,431.56.
    report = reports["monthly_business_review"]
    rows = report.numbers["sections"]["kpis_month"]["rows"]
    net = next(r for r in rows if r["kpi"] == "net_revenue")["current"]
    assert net == 462431.56
    assert "462,431.56" in report.html and "462,431.55" not in report.html


def test_week2_hand_checked_figures_unchanged(reports: dict[str, Report]) -> None:
    # docs/hand_checks/week2_weekly_2026-09-27.md
    page = visible_text(reports["weekly_sales_review"].html)
    for figure in ["114,286.82", "661", "172.90", "43.1%", "31,481.36", "2,767.45"]:
        assert figure in page, figure
    assert "Stokes, Robinson and Gough" in page


def test_prebuild_email_table_has_5_columns_and_excel_has_all() -> None:
    from src.reports import demand_sections

    frame = pd.DataFrame(
        [
            {
                "rank": 1,
                "name": "Alpha",
                "sku": "A-1",
                "priority": "urgent",
                "days_of_cover": 3.0,
                "lead_time_weeks": 4.0,
                "on_hand": 2,
                "on_order": 0,
                "recommended_units": 10,
                "unit_price": 9.99,
                "value_at_risk": 99.9,
                "rec_id": "20260930-A-1",
            }
        ]
    )
    group = demand_sections.group(
        "prebuild",
        "Pre-build",
        frame,
        demand_sections.PREBUILD_COLUMNS,
        email_columns=demand_sections.PREBUILD_EMAIL_COLUMNS,
    )
    section = {"title": "Demand alerts", "note": "", "name": "demand_alerts", "groups": [group]}
    table = demand_sections.group_views(section)[1]
    assert table["header"] == [
        "Product",
        "Days of cover",
        "Pre-build units",
        "Value at risk",
        "Reason",
    ]
    assert [c["text"] for c in table["rows"][0]] == ["Alpha", "3.0", "10", "99.90", "urgent"]
    _, sheet, _ = demand_sections.excel_frame(section)
    assert {"Id", "On hand", "Unit price", "Value at risk"} <= set(sheet.columns)
